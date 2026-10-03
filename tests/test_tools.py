"""Excel / Word / テキストツールの回帰テスト"""
import datetime
import json
import os
import signal

import openpyxl
import pytest
from docx import Document
from docx.oxml.ns import qn
from openpyxl.styles import Font

import agent
import text_tools


@pytest.fixture(autouse=True)
def workdir(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    return tmp_path


def run(name, **kwargs) -> dict:
    return json.loads(agent.execute_tool(name, kwargs))


# ── 共通 ──────────────────────────────────────────────────────────────────────

def test_bad_arguments_are_returned_as_error():
    result = run("read_cell", file_path="a.xlsx")
    assert result["success"] is False
    assert "引数" in result["error"]


def test_unknown_tool():
    assert run("no_such_tool")["success"] is False


def test_path_outside_workdir_is_rejected():
    assert run("write_cell", file_path="../x.xlsx", cell_address="A1", value=1)["success"] is False


def test_save_as_outside_workdir_is_rejected(workdir):
    run("write_cell", file_path="a.xlsx", cell_address="A1", value=1)
    assert run("save_excel", file_path="a.xlsx", save_as="../escaped.xlsx")["success"] is False
    assert not (workdir.parent / "escaped.xlsx").exists()

    run("append_paragraph", file_path="a.docx", text="x")
    assert run("save_word", file_path="a.docx", save_as="../escaped.docx")["success"] is False
    assert not (workdir.parent / "escaped.docx").exists()


def test_workdir_follows_cwd_change(workdir, monkeypatch):
    run("write_cell", file_path="a.xlsx", cell_address="A1", value=1)
    (workdir / "sub").mkdir()
    monkeypatch.chdir(workdir / "sub")
    run("write_cell", file_path="a.xlsx", cell_address="A1", value=2)
    # /cd 後の操作は移動先のファイルに対して読み書きされる
    assert openpyxl.load_workbook(workdir / "a.xlsx").active["A1"].value == 1
    assert openpyxl.load_workbook(workdir / "sub" / "a.xlsx").active["A1"].value == 2


# ── Excel ────────────────────────────────────────────────────────────────────

def test_time_value_is_serializable(workdir):
    wb = openpyxl.Workbook()
    wb.active["A1"] = datetime.time(9, 30)
    wb.save(workdir / "t.xlsx")
    result = run("read_cell", file_path="t.xlsx", cell_address="A1")
    assert result["value"] == "09:30:00"


def test_external_edit_is_not_overwritten(workdir):
    run("write_cell", file_path="c.xlsx", cell_address="A1", value="agent")
    wb = openpyxl.load_workbook(workdir / "c.xlsx")
    wb.active["B1"] = "user"
    wb.save(workdir / "c.xlsx")
    # mtime の更新を確実にする
    st = os.stat(workdir / "c.xlsx")
    os.utime(workdir / "c.xlsx", ns=(st.st_atime_ns, st.st_mtime_ns + 1_000_000))

    run("write_cell", file_path="c.xlsx", cell_address="A2", value="agent2")
    ws = openpyxl.load_workbook(workdir / "c.xlsx").active
    assert ws["B1"].value == "user"
    assert ws["A2"].value == "agent2"


def test_read_sheet_is_capped_and_compact():
    run("write_range", file_path="big.xlsx", start_cell="A1", data=[[i, i * 2] for i in range(250)])
    result = run("read_sheet", file_path="big.xlsx")
    assert len(result["rows"]) == 100
    assert result["truncated"] is True
    assert result["next_min_row"] == 101
    assert result["columns"] == ["A", "B"]
    assert result["rows"][0] == {"row": 1, "values": [0, 0]}

    tail = run("read_sheet", file_path="big.xlsx", min_row=201)
    assert tail["rows"][-1] == {"row": 250, "values": [249, 498]}
    assert "truncated" not in tail


def test_write_range_wraps_scalar_rows():
    result = run("write_range", file_path="w.xlsx", start_cell="A1", data=["abc", "de"])
    assert result["success"] is True
    assert result["range"] == "A1:A2"
    rows = run("read_sheet", file_path="w.xlsx")["rows"]
    assert [r["values"] for r in rows] == [["abc"], ["de"]]


def test_write_range_rejects_empty_data():
    assert run("write_range", file_path="w.xlsx", start_cell="A1", data=[])["success"] is False


def test_format_cell_keeps_existing_font(workdir):
    wb = openpyxl.Workbook()
    wb.active["A1"] = "x"
    wb.active["A1"].font = Font(name="Meiryo", underline="single", color="FF0000")
    wb.save(workdir / "f.xlsx")

    assert run("format_cell", file_path="f.xlsx", cell_address="A1", bold=True)["success"]
    font = openpyxl.load_workbook(workdir / "f.xlsx").active["A1"].font
    assert font.bold is True
    assert font.name == "Meiryo"
    assert font.underline == "single"
    assert font.color.rgb.endswith("FF0000")


def test_format_range_applies_to_all_cells(workdir):
    run("write_range", file_path="f.xlsx", start_cell="A1", data=[[1, 2], [3, 4]])
    result = run("format_range", file_path="f.xlsx", cell_range="A1:B2",
                 bg_color="#ffff00", horizontal_align="center")
    assert "4 セル" in result["message"]
    cell = openpyxl.load_workbook(workdir / "f.xlsx").active["B2"]
    assert cell.fill.fgColor.rgb.endswith("FFFF00")
    assert cell.alignment.horizontal == "center"


def test_add_chart_with_single_column():
    run("write_range", file_path="ch.xlsx", start_cell="A1", data=[["売上"], [1], [2], [3]])
    assert run("add_chart", file_path="ch.xlsx", chart_type="bar", data_range="A1:A4")["success"]


def test_open_excel_creates_file(workdir):
    result = run("open_excel", file_path="new.xlsx", create_if_missing=True)
    assert result["success"] is True
    assert (workdir / "new.xlsx").exists()


# ── Word ─────────────────────────────────────────────────────────────────────

def _doc_with_split_runs(path):
    doc = Document()
    para = doc.add_paragraph()
    para.add_run("Q3の")
    para.add_run("結").bold = True
    para.add_run("果です")
    doc.save(path)


def test_replace_text_across_runs(workdir):
    _doc_with_split_runs(workdir / "r.docx")
    result = run("replace_text", file_path="r.docx", old_text="Q3の結果", new_text="Q3・Q4の結果")
    assert result["replaced_count"] == 1
    assert Document(workdir / "r.docx").paragraphs[0].text == "Q3・Q4の結果です"


def test_replace_text_counts_occurrences_and_respects_limit(workdir):
    doc = Document()
    doc.add_paragraph("a-a-a")
    doc.add_paragraph("a")
    doc.save(workdir / "r.docx")

    assert run("replace_text", file_path="r.docx", old_text="a", new_text="b",
               all_occurrences=False)["replaced_count"] == 1
    assert [p.text for p in Document(workdir / "r.docx").paragraphs] == ["b-a-a", "a"]

    # 置換後の文字列に検索文字列が含まれていても無限ループしない
    assert run("replace_text", file_path="r.docx", old_text="a", new_text="aa")["replaced_count"] == 3
    assert [p.text for p in Document(workdir / "r.docx").paragraphs] == ["b-aa-aa", "aa"]


def test_replace_text_in_merged_cells_and_header(workdir):
    doc = Document()
    table = doc.add_table(rows=1, cols=2)
    merged = table.cell(0, 0).merge(table.cell(0, 1))
    merged.text = "旧社名"
    doc.sections[0].header.paragraphs[0].text = "旧社名 御中"
    doc.save(workdir / "m.docx")

    result = run("replace_text", file_path="m.docx", old_text="旧社名", new_text="新旧社名")
    # 結合セルは1回だけ置換される
    assert result["replaced_count"] == 2
    saved = Document(workdir / "m.docx")
    assert saved.tables[0].cell(0, 0).text == "新旧社名"
    assert saved.sections[0].header.paragraphs[0].text == "新旧社名 御中"


def test_format_table_replaces_shading(workdir):
    doc = Document()
    doc.add_table(rows=1, cols=1)
    doc.save(workdir / "tb.docx")
    run("format_table", file_path="tb.docx", table_index=0, row=0, col=0, bg_color="FF0000")
    run("format_table", file_path="tb.docx", table_index=0, row=0, col=0, bg_color="00FF00")
    tc_pr = Document(workdir / "tb.docx").tables[0].cell(0, 0)._tc.tcPr
    shds = tc_pr.findall(qn("w:shd"))
    assert len(shds) == 1
    assert shds[0].get(qn("w:fill")) == "00FF00"


def test_add_table_without_table_grid_style(workdir):
    # Word で作成した文書を想定して "Table Grid" スタイルを取り除く
    doc = Document()
    doc.add_paragraph("intro")
    style = doc.styles["Table Grid"]._element
    style.getparent().remove(style)
    doc.save(workdir / "t.docx")

    result = run("add_table", file_path="t.docx", data=[["a", "b"], [1, 2]])
    assert result["success"] is True
    assert len(Document(workdir / "t.docx").tables) == 1


def test_add_table_rejects_bad_index_without_side_effects(workdir):
    run("append_paragraph", file_path="t.docx", text="intro")
    assert run("add_table", file_path="t.docx", data=[["a"]], paragraph_index=5)["success"] is False
    assert run("read_document", file_path="t.docx")["tables"] == []


def test_failed_edit_does_not_leak_into_next_save(workdir):
    run("append_paragraph", file_path="d.docx", text="first")
    # 存在しないスタイルで失敗 → キャッシュ上の中途半端な段落は破棄される
    assert run("insert_paragraph", file_path="d.docx", index=0, text="ghost",
               style="No Such Style")["success"] is False
    run("append_paragraph", file_path="d.docx", text="second")
    assert [p.text for p in Document(workdir / "d.docx").paragraphs] == ["first", "second"]


def test_get_document_info_with_custom_heading_style(workdir):
    doc = Document()
    doc.add_heading("見出し", level=2)
    doc.styles.add_style("Heading Custom", 1)
    doc.add_paragraph("x", style="Heading Custom")
    doc.save(workdir / "h.docx")
    result = run("get_document_info", file_path="h.docx")
    assert result["success"] is True
    assert result["headings"] == [{"index": 0, "level": 2, "text": "見出し"}]


def test_insert_image_outside_workdir_is_rejected():
    result = run("insert_image", file_path="i.docx", image_path="/etc/hosts")
    assert result["success"] is False


# ── テキスト / Markdown ───────────────────────────────────────────────────────

@pytest.fixture
def timeout():
    def handler(*_):
        raise TimeoutError("無限ループ")
    signal.signal(signal.SIGALRM, handler)
    signal.alarm(2)
    yield
    signal.alarm(0)


def test_parse_markdown_hash_without_space(workdir, timeout):
    (workdir / "h.md").write_text("#tag\n本文\n####### seven\n")
    result = text_tools.parse_markdown("h.md")
    assert result["success"] is True
    assert result["blocks"][0]["type"] == "paragraph"


def test_parse_markdown_list_keeps_inner_numbers(workdir):
    (workdir / "l.md").write_text("- see section 3. details\n1. step 2. again\n")
    items = text_tools.parse_markdown("l.md")["blocks"][0]["items"]
    assert items == ["see section 3. details", "step 2. again"]


def test_parse_markdown_with_bom(workdir):
    (workdir / "b.md").write_bytes("﻿# タイトル\n".encode("utf-8"))
    blocks = text_tools.parse_markdown("b.md")["blocks"]
    assert blocks == [{"type": "heading", "level": 1, "text": "タイトル"}]


def test_read_text_file_shift_jis(workdir):
    (workdir / "s.txt").write_bytes("日本語テキスト".encode("cp932"))
    result = text_tools.read_text_file("s.txt")
    assert result["content"] == "日本語テキスト"
    assert result["encoding"] == "cp932"


def test_parse_inline_formatting_returns_dict():
    result = run("parse_inline_formatting", text="a **b**")
    assert result["success"] is True
    assert result["runs"][1] == {"text": "b", "bold": True, "italic": False}
